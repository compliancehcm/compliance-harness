"""Build one formatted workbook from a declarative specification.

Reads one JSON request on stdin and writes exactly one JSON envelope on stdout.
Both outcomes exit 0: `{"ok": true, ...}` when the workbook was written and
`{"ok": false, "error": ...}` when it could not be. A non-zero exit therefore
means the interpreter itself failed, which the Node half reports differently.

openpyxl, not xlsxwriter: it also reads, so a later change can amend an existing
workbook, and one library is one thing to install. The cost is that charts are
openpyxl's, which is why the chart vocabulary here is deliberately small.
"""

import json
import sys
import traceback
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

#: Column type to Excel number format. The formats are pt-BR: the separators a
#: Brazilian Excel renders from these are dot-thousands and comma-decimal.
FORMATS = {
    "text": "@",
    "number": "#,##0.00",
    "integer": "#,##0",
    "currency": 'R$ #,##0.00',
    "percent": "0.0%",
    "date": "DD/MM/YYYY",
    "month": "MM/YYYY",
}

#: Types whose values are right-aligned, because a number reads against the
#: column's right edge and text against its left.
NUMERIC_TYPES = {"number", "integer", "currency", "percent"}

#: Chart kind to the openpyxl class and, for bar, its orientation.
CHARTS = {
    "bar": (BarChart, "col"),
    "column": (BarChart, "col"),
    "hbar": (BarChart, "bar"),
    "line": (LineChart, None),
    "pie": (PieChart, None),
}


def load_rows(path):
    """Read rows from a data file, choosing the reader by extension.

    pandas is imported lazily: a workbook built from inline rows should not need
    it, and this plugin must keep working where only openpyxl is installed.
    """
    import pandas as pd

    lower = path.lower()
    if lower.endswith((".csv", ".txt")):
        frame = pd.read_csv(path)
    elif lower.endswith(".tsv"):
        frame = pd.read_csv(path, sep="\t")
    elif lower.endswith(".parquet"):
        frame = pd.read_parquet(path)
    elif lower.endswith(".json"):
        frame = pd.read_json(path)
    else:
        raise ValueError(f"cannot read {path!r}: supported sources are .csv, .tsv, .txt, .parquet, .json")
    # NaN is not JSON and not a cell value; None becomes an empty cell.
    return frame.where(frame.notna(), None).to_dict("records")


def coerce(value, column_type):
    """Coerce one cell value into what Excel should store for its column type."""
    if value is None:
        return None
    if column_type in {"date", "month"}:
        if isinstance(value, (datetime, date)):
            return value
        text = str(value).strip()
        for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%dT%H:%M:%S", "%m/%Y"):
            try:
                return datetime.strptime(text, pattern)
            except ValueError:
                continue
        # An unparseable date stays text rather than becoming a wrong date: a
        # silently shifted month is worse than a cell that is visibly not a date.
        return text
    if column_type in NUMERIC_TYPES:
        if isinstance(value, bool):
            return int(value)
        if isinstance(value, (int, float)):
            return value
        text = str(value).strip()
        if text == "":
            return None
        # Accept the Brazilian written form, since that is what a source file
        # exported from a local system carries.
        if "," in text:
            text = text.replace(".", "").replace(",", ".")
        try:
            return float(text)
        except ValueError:
            return str(value)
    return str(value)


def write_sheet(worksheet, sheet, theme, limits):
    """Write one sheet's header, body, totals and chart. Returns the row count."""
    columns = sheet["columns"]
    rows = sheet["rows"] if sheet.get("rows") is not None else load_rows(sheet["source"])

    if len(rows) > limits["maxRows"]:
        raise ValueError(
            f"sheet {sheet['name']!r} has {len(rows)} rows; at most {limits['maxRows']} are allowed"
        )

    header_fill = PatternFill("solid", fgColor=theme["headerFill"])
    header_font = Font(name=theme["fontName"], size=theme["fontSize"], bold=True, color=theme["headerFont"])
    body_font = Font(name=theme["fontName"], size=theme["fontSize"])
    band_fill = PatternFill("solid", fgColor=theme["bandFill"])
    totals_fill = PatternFill("solid", fgColor=theme["totalsFill"])
    thin = Side(style="thin", color=theme["borderColor"])
    border = Border(bottom=thin)

    for index, column in enumerate(columns, start=1):
        cell = worksheet.cell(row=1, column=index, value=column["label"])
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for offset, row in enumerate(rows):
        excel_row = offset + 2
        for index, column in enumerate(columns, start=1):
            column_type = column.get("type", "text")
            cell = worksheet.cell(row=excel_row, column=index, value=coerce(row.get(column["key"]), column_type))
            cell.font = body_font
            cell.number_format = FORMATS[column_type]
            if column_type in NUMERIC_TYPES:
                cell.alignment = Alignment(horizontal="right")
            # Banding is readability, not decoration: on a wide sheet it is what
            # keeps the eye on one record while crossing columns.
            if sheet.get("banded", True) and offset % 2 == 1:
                cell.fill = band_fill

    last_row = len(rows) + 1
    totals = sheet.get("totals") or []
    if totals and rows:
        total_row = last_row + 1
        for index, column in enumerate(columns, start=1):
            cell = worksheet.cell(row=total_row, column=index)
            cell.fill = totals_fill
            cell.font = Font(name=theme["fontName"], size=theme["fontSize"], bold=True)
            if index == 1 and column["key"] not in totals:
                cell.value = "Total"
            elif column["key"] in totals:
                letter = get_column_letter(index)
                # A formula, not a computed number: the reader can change a cell
                # and watch the total follow, which is the point of a spreadsheet.
                cell.value = f"=SUBTOTAL(109,{letter}2:{letter}{last_row})"
                cell.number_format = FORMATS[column.get("type", "number")]
                cell.alignment = Alignment(horizontal="right")

    for index, column in enumerate(columns, start=1):
        width = column.get("width")
        if width is None:
            longest = len(str(column["label"]))
            for row in rows[:200]:
                value = row.get(column["key"])
                longest = max(longest, len(str("" if value is None else value)))
            width = min(max(longest + 2, 10), 60)
        worksheet.column_dimensions[get_column_letter(index)].width = width

    if sheet.get("freezeHeader", True):
        worksheet.freeze_panes = "A2"
    if sheet.get("autofilter", True) and rows:
        worksheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{last_row}"

    chart_spec = sheet.get("chart")
    if chart_spec is not None and rows:
        add_chart(worksheet, chart_spec, columns, last_row, theme)

    return len(rows)


def add_chart(worksheet, spec, columns, last_row, theme):
    """Anchor one chart beside the data."""
    keys = [column["key"] for column in columns]
    factory, direction = CHARTS[spec["type"]]
    chart = factory()
    if direction is not None:
        chart.type = direction
    chart.title = spec.get("title")
    chart.height = 8
    chart.width = 16

    categories_index = keys.index(spec["categories"]) + 1
    categories = Reference(worksheet, min_col=categories_index, min_row=2, max_row=last_row)
    for key in spec["series"]:
        index = keys.index(key) + 1
        data = Reference(worksheet, min_col=index, min_row=1, max_row=last_row)
        chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    worksheet.add_chart(chart, f"{get_column_letter(len(columns) + 2)}2")


def run(request):
    """Build the workbook described by one request."""
    path = request["path"]
    theme = request["theme"]
    limits = request["limits"]

    workbook = Workbook()
    workbook.remove(workbook.active)

    written = []
    for sheet in request["sheets"]:
        worksheet = workbook.create_sheet(title=sheet["name"])
        count = write_sheet(worksheet, sheet, theme, limits)
        written.append({"name": sheet["name"], "rowCount": count, "columnCount": len(sheet["columns"])})

    workbook.save(path)
    return {"ok": True, "path": path, "sheets": written}


def main():
    """Read the request, run it, and emit exactly one envelope."""
    try:
        request = json.load(sys.stdin)
    except Exception as error:  # noqa: BLE001 - a malformed request is still an envelope
        json.dump({"ok": False, "error": f"the request is not JSON: {error}"}, sys.stdout)
        return

    try:
        envelope = run(request)
    except Exception as error:  # noqa: BLE001 - every build failure is the model's to fix
        last = traceback.format_exc().strip().splitlines()[-1]
        envelope = {"ok": False, "error": f"{type(error).__name__}: {error}" if str(error) else last}

    json.dump(envelope, sys.stdout, ensure_ascii=False, default=str)


if __name__ == "__main__":
    main()
